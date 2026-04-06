from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from .schemas import FinalReport


def _set_page_layout(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.70)
    section.right_margin = Inches(0.70)


def _set_default_style(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10.5)
    style.paragraph_format.space_after = Pt(4)

    for style_name, font_size, bold in [
        ("Title", 15, True),
        ("Heading 1", 12.5, True),
        ("Heading 2", 11.5, True),
    ]:
        if style_name in doc.styles:
            style_obj = doc.styles[style_name]
            style_obj.font.name = "Arial"
            style_obj.font.size = Pt(font_size)
            style_obj.font.bold = bold


def _add_title(doc: Document, report: FinalReport) -> None:
    paragraph = doc.add_paragraph(style="Title")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(report.meta.title)
    run.bold = True

    meta = doc.add_paragraph()
    meta.add_run("Tanggal Pengukuran: ").bold = True
    meta.add_run(report.meta.tanggal_pengukuran)
    meta.add_run("\nArea: ").bold = True
    meta.add_run(report.meta.area)
    meta.add_run("\nFokus Laporan: ").bold = True
    meta.add_run(report.meta.fokus_laporan)


def _add_standard_reference(doc: Document, report: FinalReport) -> None:
    doc.add_paragraph("Standar Acuan Analisa:", style="Heading 1")
    doc.add_paragraph(report.standard_reference_intro)
    for item in report.thresholds:
        doc.add_paragraph(item)


def _format_table(table) -> None:
    widths = [0.5, 1.0, 1.8, 1.9, 3.1, 2.4]

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.space_before = Pt(1)
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(9.5)

    for cell in table.rows[0].cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True


def _add_analysis_table(doc: Document, report: FinalReport) -> None:
    doc.add_paragraph("Tabel Analisa & Rekomendasi", style="Heading 1")

    table = doc.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "No."
    hdr[1].text = "Tag Name"
    hdr[2].text = "Deskripsi"
    hdr[3].text = "Titik & Nilai Vibrasi Tertinggi"
    hdr[4].text = "Analisa"
    hdr[5].text = "Rekomendasi"

    for row in report.rows:
        cells = table.add_row().cells
        cells[0].text = str(row.no)
        cells[1].text = row.tag_name
        cells[2].text = row.deskripsi
        cells[3].text = row.titik_nilai_vibrasi_tertinggi
        cells[4].text = row.analysis
        cells[5].text = "\n".join(row.recommendations)

    _format_table(table)


def _add_priority_summary(doc: Document, report: FinalReport) -> None:
    doc.add_paragraph("Ringkasan dan Prioritas Tindakan (Terkoreksi)", style="Heading 1")

    p1 = doc.add_paragraph()
    p1.add_run("1. Prioritas 1: Tindakan Segera (Kategori BAHAYA)\n").bold = True
    p1.add_run(report.priority_1_text)

    p2 = doc.add_paragraph()
    p2.add_run("2. Prioritas 2: Jadwalkan Investigasi (Kategori WASPADA)\n").bold = True
    p2.add_run(report.priority_2_text)

    p3 = doc.add_paragraph()
    p3.add_run("3. Prioritas 3: Monitoring Diperketat (Kategori CUKUP / BAIK)\n").bold = True
    p3.add_run(report.priority_3_text)

    doc.add_paragraph(report.closing_note)


def write_report_docx(report: FinalReport, output_path: str | Path) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    _set_page_layout(doc)
    _set_default_style(doc)
    _add_title(doc, report)
    doc.add_paragraph("")
    _add_standard_reference(doc, report)
    doc.add_paragraph("")
    _add_analysis_table(doc, report)
    doc.add_paragraph("")
    _add_priority_summary(doc, report)

    doc.save(str(output_path))
    return output_path