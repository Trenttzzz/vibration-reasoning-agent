from __future__ import annotations

import json
from pathlib import Path
from typing import Any

import openpyxl


COLUMN_MAP = {
    3: "no",
    4: "area",
    5: "equipment.tag_name",
    6: "equipment.deskripsi",
    7: "equipment.pondasi",
    8: "equipment.posisi",
    9: "equipment.kw",
    10: "equipment.ampere",
    11: "equipment.rpm",
    12: "status.run",
    13: "status.standby",
    14: "status.repair",
    15: "driver.NDE.ax",
    16: "driver.NDE.v",
    17: "driver.NDE.h",
    18: "driver.NDE.he",
    19: "driver.NDE.temp",
    20: "driver.DE.ax",
    21: "driver.DE.v",
    22: "driver.DE.h",
    23: "driver.DE.he",
    24: "driver.DE.temp",
    25: "driver.grounding",
    26: "driver.temp_terminal_box",
    27: "driven.DE.ax",
    28: "driven.DE.v",
    29: "driven.DE.h",
    30: "driven.DE.he",
    31: "driven.DE.temp",
    32: "driven.NDE.ax",
    33: "driven.NDE.v",
    34: "driven.NDE.h",
    35: "driven.NDE.he",
    36: "driven.NDE.temp",
}

DATA_START_ROW = 12
STATUS_MARK = "√"


def _set_nested(d: dict, dotted_key: str, value: Any) -> None:
    keys = dotted_key.split(".")
    for key in keys[:-1]:
        d = d.setdefault(key, {})
    d[keys[-1]] = value


def _parse_status(value: Any) -> bool:
    if value is None:
        return False
    return str(value).strip() == STATUS_MARK


def _parse_numeric(value: Any) -> float | None:
    if value is None:
        return None
    try:
        return float(value)
    except (ValueError, TypeError):
        return None


def _parse_cell(col: int, value: Any) -> Any:
    key = COLUMN_MAP.get(col, "")

    if key.startswith("status."):
        return _parse_status(value)

    if any(
        key.endswith(suffix)
        for suffix in (".ax", ".v", ".h", ".he", ".temp", ".kw", ".ampere", ".rpm")
    ):
        return _parse_numeric(value)

    if value is None:
        return None

    return str(value).strip()


def _get_sheet_date(ws) -> str | None:
    for row in ws.iter_rows(min_row=7, max_row=7, min_col=4, max_col=6):
        for cell in row:
            if cell.value and str(cell.value).strip().startswith("Date"):
                date_cell = ws.cell(row=7, column=cell.column + 1)
                if date_cell.value:
                    return str(date_cell.value).strip()
    return ws.title


def _ensure_keys(record: dict) -> None:
    record.setdefault("area", None)
    record.setdefault("equipment", {})
    record.setdefault("status", {"run": False, "standby": False, "repair": False})
    record.setdefault("driver", {})
    record.setdefault("driven", {})

    for bearing in ("NDE", "DE"):
        record["driver"].setdefault(
            bearing,
            {"ax": None, "v": None, "h": None, "he": None, "temp": None},
        )
        record["driven"].setdefault(
            bearing,
            {"ax": None, "v": None, "h": None, "he": None, "temp": None},
        )

    record["driver"].setdefault("grounding", None)
    record["driver"].setdefault("temp_terminal_box", None)


def parse_sheet(ws) -> list[dict]:
    date_str = _get_sheet_date(ws)
    results: list[dict] = []

    for row_idx in range(DATA_START_ROW, ws.max_row + 1):
        no_cell = ws.cell(row=row_idx, column=3)
        area_cell = ws.cell(row=row_idx, column=4)
        tag_cell = ws.cell(row=row_idx, column=5)
        desc_cell = ws.cell(row=row_idx, column=6)

        no_val = no_cell.value
        area_val = area_cell.value
        tag_val = tag_cell.value
        desc_val = desc_cell.value

        # Skip row kosong total
        if not any([no_val, area_val, tag_val, desc_val]):
            continue

        # Skip row yang bukan data equipment valid
        if no_val is None:
            continue
        if tag_val is None or str(tag_val).strip() == "":
            continue
        if desc_val is None or str(desc_val).strip() == "":
            continue
        if area_val is None or str(area_val).strip() == "":
            continue

        record: dict = {"tanggal": date_str}

        for col, key in COLUMN_MAP.items():
            raw = ws.cell(row=row_idx, column=col).value
            parsed = _parse_cell(col, raw)
            if parsed is not None:
                _set_nested(record, key, parsed)

        _ensure_keys(record)
        results.append(record)

    return results


def parse_excel_file(filepath: str | Path) -> list[dict]:
    wb = openpyxl.load_workbook(filepath, data_only=True)
    all_records: list[dict] = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        records = parse_sheet(ws)
        all_records.extend(records)
    return all_records


def save_parsed_json(records: list[dict], output_path: str | Path) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with open(output_path, "w", encoding="utf-8") as file:
        json.dump(records, file, ensure_ascii=False, indent=2)
    return output_path


def _vib_str(bearing: dict, label: str) -> str:
    parts = []
    for key in ("ax", "v", "h", "he"):
        value = bearing.get(key)
        if value is not None:
            parts.append(f"{key.upper()}={value}")
    temp = bearing.get("temp")
    temp_str = f", Temp={temp}°C" if temp is not None else ""
    vib_str = "/".join(parts) if parts else "no data"
    return f"{label}: vib={vib_str}{temp_str}"


def format_for_llm(equipment: dict) -> str:
    eq = equipment.get("equipment", {})
    stat = equipment.get("status", {})
    drv = equipment.get("driver", {})
    drn = equipment.get("driven", {})

    status_parts = [key for key, value in stat.items() if value]
    status_str = status_parts[0].upper() if status_parts else "UNKNOWN"

    lines = [
        f"[Equipment #{equipment.get('no')}]",
        f"Tanggal  : {equipment.get('tanggal')}",
        f"Area     : {equipment.get('area')}",
        f"Tag      : {eq.get('tag_name')} — {eq.get('deskripsi')}",
        (
            f"Spec     : {eq.get('pondasi')} | {eq.get('posisi')} | "
            f"KW={eq.get('kw')} A={eq.get('ampere')} RPM={eq.get('rpm')}"
        ),
        f"Status   : {status_str}",
        "",
        "DRIVER (Motor):",
        f"  {_vib_str(drv.get('NDE', {}), 'NDE')}",
        f"  {_vib_str(drv.get('DE', {}), 'DE')}",
        (
            f"  Grounding={drv.get('grounding')} | "
            f"Temp.Terminal Box={drv.get('temp_terminal_box')}°C"
            if drv.get("temp_terminal_box") is not None
            else f"  Grounding={drv.get('grounding')}"
        ),
        "",
        "DRIVEN (Gearbox/Pump/Fan/etc):",
        f"  {_vib_str(drn.get('DE', {}), 'DE')}",
        f"  {_vib_str(drn.get('NDE', {}), 'NDE')}",
    ]
    return "\n".join(lines)


if __name__ == "__main__":
    import sys

    default_path = Path("../input/template.xlsx")
    path = sys.argv[1] if len(sys.argv) > 1 else str(default_path)

    records = parse_excel_file(path)
    print(f"Parsed {len(records)} equipment records")

    preview_count = min(3, len(records))
    for rec in records[:preview_count]:
        print("=" * 70)
        print(format_for_llm(rec))

    output_json = Path(path).with_suffix("").name + "_parsed.json"
    save_parsed_json(records, output_json)
    print(f"Saved parsed JSON to: {output_json}")